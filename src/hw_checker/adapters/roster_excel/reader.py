"""Excel (.xlsx) implementation of the RosterReader port, using openpyxl."""

import zipfile
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path

from openpyxl import load_workbook
from openpyxl.cell.cell import Cell, MergedCell
from openpyxl.utils.exceptions import InvalidFileException
from openpyxl.workbook.workbook import Workbook
from openpyxl.worksheet.worksheet import Worksheet

from hw_checker.domain import RosterRow
from hw_checker.domain.errors import RosterError

XLSX_SUFFIX = ".xlsx"


@dataclass(frozen=True)
class RosterColumns:
    """Header texts of the roster columns, matched after trimming and casefolding."""

    name: str
    partner_name: str
    github_url: str


class ExcelRosterReader:
    """Reads roster rows from one sheet of an `.xlsx` file, finding columns by header text.

    Exceptions from openpyxl and the file system are mapped to `RosterError`.
    """

    def __init__(self, columns: RosterColumns, sheet_name: str | None, header_row: int) -> None:
        self._columns = columns
        self._sheet_name = sheet_name
        self._header_row = header_row

    def read_rows(self, source: Path) -> tuple[RosterRow, ...]:
        """Return every non-empty data row below the header row, with its Excel row number.

        Raises:
            RosterError: the file is missing, not a valid `.xlsx`, or lacks a required
                sheet, header or student name.
        """
        workbook = _open_workbook(source)
        try:
            sheet = self._select_sheet(workbook, source)
            positions = self._find_columns(sheet, source)
            return tuple(self._read_data_rows(sheet, positions, source))
        finally:
            workbook.close()

    def _select_sheet(self, workbook: Workbook, source: Path) -> Worksheet:
        if self._sheet_name is None:
            sheet = workbook.worksheets[0]
        elif self._sheet_name in workbook.sheetnames:
            sheet = workbook[self._sheet_name]
        else:
            raise RosterError(
                f"Roster {source}: sheet '{self._sheet_name}' not found. "
                f"Sheets in the file: {', '.join(workbook.sheetnames)}."
            )
        if not isinstance(sheet, Worksheet):
            raise RosterError(f"Roster {source}: '{sheet.title}' is not a worksheet.")
        return sheet

    def _find_columns(self, sheet: Worksheet, source: Path) -> dict[str, int]:
        found: dict[str, int] = {}
        headers: list[str] = []
        for column, cell in enumerate(sheet[self._header_row], start=1):
            text = _cell_text(cell)
            if text:
                headers.append(text)
                found.setdefault(_normalize_header(text), column)
        positions: dict[str, int] = {}
        for field, header in vars(self._columns).items():
            position = found.get(_normalize_header(header))
            if position is None:
                listed = ", ".join(f"'{text}'" for text in headers) or "none"
                raise RosterError(
                    f"Roster {source}: header '{header}' not found in row {self._header_row} "
                    f"of sheet '{sheet.title}'. Headers found: {listed}."
                )
            positions[field] = position
        return positions

    def _read_data_rows(
        self, sheet: Worksheet, positions: dict[str, int], source: Path
    ) -> list[RosterRow]:
        rows: list[RosterRow] = []
        for row_number in range(self._header_row + 1, sheet.max_row + 1):
            name = _cell_text(sheet.cell(row_number, positions["name"]))
            partner = _cell_text(sheet.cell(row_number, positions["partner_name"]))
            url = _link_or_text(sheet.cell(row_number, positions["github_url"]))
            if not (name or partner or url):
                continue
            if not name:
                raise RosterError(
                    f"Roster {source}: row {row_number} has no student name. "
                    "Fill in the name or delete the row."
                )
            rows.append(
                RosterRow(
                    row_number=row_number,
                    name=name,
                    partner_name=partner or None,
                    github_url=url,
                )
            )
        return rows


def _open_workbook(source: Path) -> Workbook:
    if source.suffix.lower() != XLSX_SUFFIX:
        raise RosterError(
            f"Roster {source}: only .xlsx files are supported. "
            "Open the file in Excel and save it as .xlsx."
        )
    try:
        # Not read-only: hyperlinks are only available on fully loaded worksheets.
        return load_workbook(source, data_only=True)
    except FileNotFoundError as error:
        raise RosterError(f"Roster {source}: file not found.") from error
    except (InvalidFileException, zipfile.BadZipFile, KeyError, ValueError) as error:
        raise RosterError(f"Roster {source}: not a valid .xlsx file ({error}).") from error
    except OSError as error:
        raise RosterError(f"Roster {source}: cannot read the file ({error}).") from error


def _normalize_header(text: str) -> str:
    return " ".join(text.split()).casefold()


def _link_or_text(cell: Cell | MergedCell) -> str:
    hyperlink = getattr(cell, "hyperlink", None)
    target = getattr(hyperlink, "target", None)
    if isinstance(target, str) and target.strip():
        return target.strip()
    return _cell_text(cell)


def _cell_text(cell: Cell | MergedCell) -> str:
    value = cell.value
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, datetime | date | time):
        return value.isoformat()
    return str(value).strip()
