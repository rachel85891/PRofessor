from collections.abc import Sequence
from datetime import date
from pathlib import Path

import pytest
from openpyxl import Workbook, load_workbook

from hw_checker.adapters.roster_excel import ExcelRosterReader, RosterColumns
from hw_checker.domain import RosterRow
from hw_checker.domain.errors import RosterError
from hw_checker.ports import RosterReader

COLUMNS = RosterColumns(name="Name", partner_name="Partner", github_url="GitHub")
HEADER = ["Name", "Partner", "GitHub"]
DANA_URL = "https://github.com/dana/api"


def reader(sheet_name: str | None = None, header_row: int = 1) -> ExcelRosterReader:
    return ExcelRosterReader(COLUMNS, sheet_name=sheet_name, header_row=header_row)


def write_xlsx(path: Path, rows: Sequence[Sequence[object]], title: str = "Roster") -> Path:
    workbook = Workbook()
    sheet = workbook.active
    assert sheet is not None
    sheet.title = title
    for values in rows:
        sheet.append(list(values))
    workbook.save(path)
    return path


def test_reader_satisfies_the_port() -> None:
    port: RosterReader = reader()

    assert isinstance(port, ExcelRosterReader)


def test_reads_rows_with_excel_row_numbers(tmp_path: Path) -> None:
    source = write_xlsx(
        tmp_path / "roster.xlsx",
        [
            HEADER,
            ["Dana Levi", "Noa Cohen", DANA_URL],
            ["Noa Cohen", None, "https://github.com/noa/api"],
        ],
    )

    rows = reader().read_rows(source)

    assert rows == (
        RosterRow(row_number=2, name="Dana Levi", partner_name="Noa Cohen", github_url=DANA_URL),
        RosterRow(row_number=3, name="Noa Cohen", github_url="https://github.com/noa/api"),
    )


def test_columns_are_found_in_any_order_among_other_columns(tmp_path: Path) -> None:
    source = write_xlsx(
        tmp_path / "roster.xlsx",
        [["GitHub", "Notes", "Name", "Partner"], [DANA_URL, "late", "Dana", "Noa"]],
    )

    (row,) = reader().read_rows(source)

    assert (row.name, row.partner_name, row.github_url) == ("Dana", "Noa", DANA_URL)


def test_header_matching_ignores_case_and_whitespace(tmp_path: Path) -> None:
    source = write_xlsx(
        tmp_path / "roster.xlsx", [["  NAME ", "partner", " github  "], ["Dana", None, DANA_URL]]
    )

    (row,) = reader().read_rows(source)

    assert row.github_url == DANA_URL


def test_missing_header_names_it_and_lists_found_headers(tmp_path: Path) -> None:
    source = write_xlsx(tmp_path / "roster.xlsx", [["Name", "Repo"], ["Dana", DANA_URL]])

    with pytest.raises(RosterError) as raised:
        reader().read_rows(source)

    message = str(raised.value)
    assert "header 'Partner' not found" in message
    assert "'Name', 'Repo'" in message


def test_hyperlink_target_wins_over_display_text(tmp_path: Path) -> None:
    source = write_xlsx(tmp_path / "roster.xlsx", [HEADER, ["Dana", None, "my repo"]])
    workbook = load_workbook(source)
    sheet = workbook.active
    assert sheet is not None
    sheet.cell(2, 3).hyperlink = "https://github.com/dana/real-repo"
    workbook.save(source)

    (row,) = reader().read_rows(source)

    assert row.github_url == "https://github.com/dana/real-repo"


def test_empty_rows_are_skipped_and_row_numbers_kept(tmp_path: Path) -> None:
    source = write_xlsx(
        tmp_path / "roster.xlsx",
        [HEADER, ["Dana", None, DANA_URL], [None, None, None], ["Noa", None, "x"]],
    )

    rows = reader().read_rows(source)

    assert [row.row_number for row in rows] == [2, 4]


def test_row_with_name_but_no_url_is_returned(tmp_path: Path) -> None:
    source = write_xlsx(tmp_path / "roster.xlsx", [HEADER, ["Dana", "Noa", None]])

    (row,) = reader().read_rows(source)

    assert row.github_url == ""


def test_row_with_url_but_no_name_raises_naming_the_row(tmp_path: Path) -> None:
    source = write_xlsx(
        tmp_path / "roster.xlsx", [HEADER, ["Dana", None, DANA_URL], [None, None, "x"]]
    )

    with pytest.raises(RosterError, match="row 3 has no student name"):
        reader().read_rows(source)


def test_non_text_cells_become_trimmed_text(tmp_path: Path) -> None:
    source = write_xlsx(
        tmp_path / "roster.xlsx",
        [HEADER, [12345, 7.0, "  " + DANA_URL + "  "], [date(2026, 10, 20), 2.5, "x"]],
    )

    first, second = reader().read_rows(source)

    assert (first.name, first.partner_name, first.github_url) == ("12345", "7", DANA_URL)
    assert second.name.startswith("2026-10-20")
    assert second.partner_name == "2.5"


def test_non_default_sheet_and_header_row(tmp_path: Path) -> None:
    workbook = Workbook()
    first = workbook.active
    assert first is not None
    first.title = "Cover"
    first.append(["This sheet is not the roster"])
    sheet = workbook.create_sheet("Students")
    sheet.append(["Course roster"])
    sheet.append([])
    sheet.append(HEADER)
    sheet.append(["Dana", None, DANA_URL])
    source = tmp_path / "roster.xlsx"
    workbook.save(source)

    (row,) = reader(sheet_name="Students", header_row=3).read_rows(source)

    assert (row.row_number, row.name) == (4, "Dana")


def test_unknown_sheet_lists_the_sheets(tmp_path: Path) -> None:
    source = write_xlsx(tmp_path / "roster.xlsx", [HEADER], title="Roster")

    with pytest.raises(RosterError, match=r"sheet 'Students' not found.*Roster"):
        reader(sheet_name="Students").read_rows(source)


def test_missing_file_raises_roster_error(tmp_path: Path) -> None:
    with pytest.raises(RosterError, match="file not found"):
        reader().read_rows(tmp_path / "absent.xlsx")


def test_text_file_named_xlsx_raises_roster_error(tmp_path: Path) -> None:
    source = tmp_path / "roster.xlsx"
    source.write_text("name,partner,github\n", encoding="utf-8")

    with pytest.raises(RosterError, match=r"not a valid \.xlsx file"):
        reader().read_rows(source)


@pytest.mark.parametrize("name", ["roster.xls", "roster.csv"])
def test_other_extensions_raise_roster_error(tmp_path: Path, name: str) -> None:
    source = tmp_path / name
    source.write_bytes(b"\xd0\xcf\x11\xe0")

    with pytest.raises(RosterError, match=r"only \.xlsx files are supported"):
        reader().read_rows(source)
